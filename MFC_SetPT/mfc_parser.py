import os
import csv
from datetime import datetime
import tkinter as tk
from tkinter import filedialog, messagebox
from tkinterdnd2 import TkinterDnD, DND_FILES
import openpyxl
# 엑셀 꾸미기를 위한 스타일 모듈 추가
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

selected_files = []

# ================================
# 도움말(Help) 메시지 함수
# ================================
def show_general_help():
    help_text = """[ MFC 유량 데이터 추출기 사용법 ]

1. [파일 추가] 
   드래그 앤 드롭 또는 'Browse' 버튼으로 분석할 CSV 파일들을 추가합니다.
   
2. [옵션 설정] 
   하단의 저장 설정에서 원하는 파일명 형태와 확장자(Excel/CSV)를 선택합니다.
   (시간 자동 변환 기능을 켜고 끌 수 있습니다. [?] 버튼을 눌러 규칙을 확인하세요.)
   
3. [변환 실행] 
   'Process & Save' 버튼을 누르면, 전체 파일의 데이터가 시간순으로 통합되어 하나의 파일로 깔끔하게 정리됩니다.

* 팁: 잘못 추가한 파일은 목록에서 클릭 후 'Remove'를 누르면 지워집니다.

--------------------------------------------------
System Designer : 한국공학대학교 메카트로닉스공학부 21학번 황영진
"""
    messagebox.showinfo("프로그램 사용법", help_text)

def show_time_format_help():
    help_text = """[ 파일명 시간 자동변환 예약어 ]

'시간 자동 변환 적용' 체크박스가 켜져 있을 때,
아래의 영문을 파일명에 입력하면 현재 시간으로 변환됩니다.

• yyyy : 연도 4자리 (예: 2026)
• yy : 연도 2자리 (예: 26)
• MM : 월 2자리 (예: 09)
• dd : 일 2자리 (예: 17)
• hh 또는 HH : 시간 2자리 (예: 14)
• mm : 분 2자리 (예: 20)
• ss : 초 2자리 (예: 03)

▶ 사용 예시
입력: yyMMdd_hhmm_mfc_setpt
결과: 260917_1420_mfc_setpt

※ 윈도우 파일명에는 콜론(:)을 쓸 수 없으므로, hh:mm 입력 시 hh-mm으로 자동 변환됩니다.
※ 영문 파일명(summary 등)을 쓸 때 mm이나 ss가 숫자로 바뀌는 것을 원치 않으시면 '시간 자동 변환 적용' 체크박스를 해제해 주세요."""
    messagebox.showinfo("시간 형식 설명", help_text)

# ================================
# 파일명 자동 포맷팅 함수
# ================================
def get_formatted_filename(pattern):
    now = datetime.now()
    res = pattern
    
    res = res.replace("yyyy", now.strftime("%Y"))
    res = res.replace("yy", now.strftime("%y"))
    res = res.replace("MM", now.strftime("%m"))
    res = res.replace("dd", now.strftime("%d"))
    res = res.replace("HH", now.strftime("%H"))
    res = res.replace("hh", now.strftime("%H"))
    res = res.replace("mm", now.strftime("%M"))
    res = res.replace("ss", now.strftime("%S"))
    
    invalid_chars = ['\\', '/', ':', '*', '?', '"', '<', '>', '|']
    for char in invalid_chars:
        if char == ':':
            res = res.replace(char, "-")
        else:
            res = res.replace(char, "_")
            
    return res

# ================================
# 핵심 기능 함수
# ================================
def update_listbox():
    listbox.delete(0, tk.END)
    for path in selected_files:
        listbox.insert(tk.END, os.path.basename(path))

def select_files():
    filepaths = filedialog.askopenfilenames(
        title="Select MFC CSV Files",
        filetypes=(("CSV files", "*.csv"), ("All files", "*.*"))
    )
    for path in filepaths:
        if path not in selected_files:
            selected_files.append(path)
    update_listbox()

def drop_files(event):
    files = root.tk.splitlist(event.data)
    for path in files:
        if path.lower().endswith('.csv') and path not in selected_files:
            selected_files.append(path)
    update_listbox()

def remove_selected():
    selected_indices = listbox.curselection()
    for index in reversed(selected_indices):
        del selected_files[index]
    update_listbox()
    
def clear_files():
    selected_files.clear()
    update_listbox()

def process_and_save():
    if not selected_files:
        messagebox.showwarning("Warning", "먼저 CSV 파일을 선택하거나 드래그 앤 드롭 해주세요.")
        return

    raw_pattern = entry_filename.get()
    if use_time_format_var.get():
        formatted_name = get_formatted_filename(raw_pattern)
    else:
        formatted_name = raw_pattern 

    ext = ext_var.get()
    default_full_name = f"{formatted_name}{ext}"

    save_path = filedialog.asksaveasfilename(
        title="결과 저장",
        defaultextension=ext,
        initialfile=default_full_name,
        filetypes=((f"{ext.upper()} files", f"*{ext}"), ("All files", "*.*"))
    )

    if not save_path:
        return

    try:
        results = []
        for filepath in selected_files:
            filename = os.path.basename(filepath)
            
            with open(filepath, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                rows = list(reader)

                if not rows: continue

                ar_setpts = [float(row['Ar_SetPT']) for row in rows]
                n2_setpts = [float(row['N2_SetPT']) for row in rows]

                ar_is_bg = (min(ar_setpts) == max(ar_setpts))
                n2_is_bg = (min(n2_setpts) == max(n2_setpts))

                if ar_is_bg and not n2_is_bg:
                    bg_gas, ctrl_gas = 'Ar', 'N2'
                    bg_col, ctrl_col = 'Ar_SetPT', 'N2_SetPT'
                elif n2_is_bg and not ar_is_bg:
                    bg_gas, ctrl_gas = 'N2', 'Ar'
                    bg_col, ctrl_col = 'N2_SetPT', 'Ar_SetPT'
                else:
                    continue

                bg_flow = float(rows[0][bg_col])
                prev_ctrl_flow = float(rows[0][ctrl_col])
                bg_flow_str = int(bg_flow) if bg_flow.is_integer() else bg_flow

                for i in range(1, len(rows)):
                    curr_ctrl_flow = float(rows[i][ctrl_col])

                    if curr_ctrl_flow != prev_ctrl_flow:
                        change_time = rows[i-1]['Time'] 
                        prev_str = int(prev_ctrl_flow) if prev_ctrl_flow.is_integer() else prev_ctrl_flow
                        curr_str = int(curr_ctrl_flow) if curr_ctrl_flow.is_integer() else curr_ctrl_flow

                        results.append([
                            filename, change_time, bg_gas, bg_flow_str, ctrl_gas, prev_str, curr_str
                        ])
                        prev_ctrl_flow = curr_ctrl_flow

        results.sort(key=lambda x: x[1])
        headers = ['Filename', 'Time', 'BG_Gas', 'BG_Flow', 'Ctrl_Gas', 'Prev_Flow', 'Curr_Flow']

        if save_path.endswith('.xlsx'):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "MFC Data"
            
            # 스타일 정의
            # 헤더: 짙은 파란색 배경, 흰색 볼드 폰트
            header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            # 정렬: 가운데 정렬
            center_align = Alignment(horizontal="center", vertical="center")

            # 1. 헤더 추가 및 서식 적용
            ws.append(headers)
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align

            # 2. 데이터 추가 및 정렬 서식 적용
            for row_idx, row_data in enumerate(results, start=2):
                ws.append(row_data)
                for col_idx in range(1, len(row_data) + 1):
                    ws.cell(row=row_idx, column=col_idx).alignment = center_align

            # 3. 글자 길이에 맞춰 열 너비 자동 조절
            for col in ws.columns:
                max_length = 0
                column_letter = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                # 여백을 위해 약간(+2) 더 넓게 설정
                ws.column_dimensions[column_letter].width = max_length + 2

            wb.save(save_path)
        else:
            with open(save_path, 'w', newline='', encoding='utf-8-sig') as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                writer.writerows(results)

        messagebox.showinfo("Success", "데이터 정리가 성공적으로 완료되었습니다!")

    except Exception as e:
        messagebox.showerror("Error", f"저장 중 에러가 발생했습니다:\n{e}")

# ================================
# GUI 화면 구성
# ================================
root = TkinterDnD.Tk()
root.title("MFC Flow Extractor")
root.geometry("540x530")

top_frame = tk.Frame(root)
top_frame.pack(fill="x", pady=5, padx=20)

lbl = tk.Label(top_frame, text="Drag & Drop CSV files here or use 'Browse'.", font=("", 10))
lbl.pack(side="left")

btn_help = tk.Button(top_frame, text="사용법 (?)", command=show_general_help, bg="#e0e0e0")
btn_help.pack(side="right")

listbox = tk.Listbox(root, selectmode=tk.MULTIPLE, width=70, height=10)
listbox.pack(pady=5)
listbox.drop_target_register(DND_FILES)
listbox.dnd_bind('<<Drop>>', drop_files)

btn_frame = tk.Frame(root)
btn_frame.pack(pady=5)

btn_select = tk.Button(btn_frame, text="Browse", command=select_files, width=10)
btn_select.grid(row=0, column=0, padx=5)
btn_remove = tk.Button(btn_frame, text="Remove", command=remove_selected, width=10)
btn_remove.grid(row=0, column=1, padx=5)
btn_clear = tk.Button(btn_frame, text="Clear All", command=clear_files, width=10)
btn_clear.grid(row=0, column=2, padx=5)

setting_frame = tk.LabelFrame(root, text="저장 설정", padx=10, pady=10)
setting_frame.pack(pady=10, fill="x", padx=20)

use_time_format_var = tk.BooleanVar(value=True)
chk_time_format = tk.Checkbutton(setting_frame, text="시간 자동 변환 적용 (yy, MM, mm 등)", variable=use_time_format_var)
chk_time_format.grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 5))

lbl_filename = tk.Label(setting_frame, text="파일명 형식:")
lbl_filename.grid(row=1, column=0, sticky="w", pady=5)

entry_filename = tk.Entry(setting_frame, width=32)
entry_filename.insert(0, "yyMMdd_hhmm_mfc_setpt") 
entry_filename.grid(row=1, column=1, sticky="w", padx=5)

btn_time_help = tk.Button(setting_frame, text="?", command=show_time_format_help, width=2, bg="#e0e0e0")
btn_time_help.grid(row=1, column=2, sticky="w")

lbl_ext = tk.Label(setting_frame, text="확장자:")
lbl_ext.grid(row=2, column=0, sticky="w", pady=5)

ext_var = tk.StringVar(value=".xlsx") 
ext_frame = tk.Frame(setting_frame)
ext_frame.grid(row=2, column=1, columnspan=2, sticky="w")
rb_xlsx = tk.Radiobutton(ext_frame, text="Excel (.xlsx)", variable=ext_var, value=".xlsx")
rb_xlsx.pack(side="left")
rb_csv = tk.Radiobutton(ext_frame, text="CSV (.csv)", variable=ext_var, value=".csv")
rb_csv.pack(side="left")

btn_process = tk.Button(root, text="Process & Save", command=process_and_save, width=20, height=2, bg="#4CAF50", fg="white", font=("", 10, "bold"))
btn_process.pack(pady=5)

root.mainloop()