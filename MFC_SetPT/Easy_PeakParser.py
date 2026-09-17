import os
import csv
from datetime import datetime, timedelta
import tkinter as tk
from tkinter import filedialog, messagebox
from tkinterdnd2 import TkinterDnD, DND_FILES
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

mfc_files = []
peak_files = []

# ================================
# 시간 포맷 정밀 파싱 함수
# ================================
def parse_mfc_time(t_str):
    formats = [
        "%y%m%d %H:%M:%S.%f",
        "%y%m%d %H:%M:%S",
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S"
    ]
    t_str = str(t_str).strip()
    for fmt in formats:
        try:
            return datetime.strptime(t_str, fmt)
        except ValueError:
            continue
    return None

def parse_peak_time(t_str):
    try:
        parts = str(t_str).strip().split(':')
        if len(parts) >= 7:
            ms = parts[6]
            if len(ms) == 3: ms += "000"
            t_str_fixed = ":".join(parts[:6]) + "." + ms
            return datetime.strptime(t_str_fixed, "%Y:%m:%d:%H:%M:%S.%f")
    except:
        pass
    return None

# [수정됨] base_time 매개변수를 추가하여 현재 시간이 아닌 '측정 시간'을 기준으로 변환
def get_formatted_filename(pattern, base_time=None):
    if base_time is None:
        base_time = datetime.now()
        
    res = pattern
    res = res.replace("yyyy", base_time.strftime("%Y")).replace("yy", base_time.strftime("%y"))
    res = res.replace("MM", base_time.strftime("%m")).replace("dd", base_time.strftime("%d"))
    res = res.replace("HH", base_time.strftime("%H")).replace("hh", base_time.strftime("%H"))
    res = res.replace("mm", base_time.strftime("%M")).replace("ss", base_time.strftime("%S"))
    for char in ['\\', '/', ':', '*', '?', '"', '<', '>', '|']:
        res = res.replace(char, "-" if char == ':' else "_")
    return res

# ================================
# MFC 데이터 추출기 (Raw & 정리본 겸용)
# ================================
def extract_mfc_events(filepath):
    events = []
    ext = filepath.lower()
    
    if ext.endswith('.csv'):
        with open(filepath, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            rows = list(reader)
            if not rows: return events
            
            if 'BG_Gas' in reader.fieldnames and 'Prev_Flow' in reader.fieldnames:
                for row in rows:
                    dt = parse_mfc_time(row['Time'])
                    if dt:
                        events.append({
                            'time': dt, 'bg_gas': row['BG_Gas'], 'bg_flow': float(row['BG_Flow']),
                            'ctrl_gas': row['Ctrl_Gas'], 'prev_flow': float(row['Prev_Flow']), 'curr_flow': float(row['Curr_Flow'])
                        })
            elif 'Ar_SetPT' in reader.fieldnames: 
                ar_setpts = [float(row['Ar_SetPT']) for row in rows]
                n2_setpts = [float(row['N2_SetPT']) for row in rows]
                ar_is_bg = (min(ar_setpts) == max(ar_setpts))
                n2_is_bg = (min(n2_setpts) == max(n2_setpts))

                if ar_is_bg and not n2_is_bg:
                    bg_gas, ctrl_gas = 'Ar', 'N2'
                    bg_col, ctrl_col = 'Ar_SetPT', 'N2_SetPT'
                elif n2_is_bg and not ar_is_bg:
                    bg_gas, ctrl_gas = 'N2', 'Ar'
                    bg_col, ctrl_col = 'N2_SetPT', 'Ar_SetPT'
                else:
                    return events

                bg_flow = float(rows[0][bg_col])
                prev_ctrl = float(rows[0][ctrl_col])
                for i in range(1, len(rows)):
                    curr_ctrl = float(rows[i][ctrl_col])
                    if curr_ctrl != prev_ctrl:
                        dt = parse_mfc_time(rows[i-1]['Time'])
                        if dt:
                            events.append({
                                'time': dt, 'bg_gas': bg_gas, 'bg_flow': bg_flow,
                                'ctrl_gas': ctrl_gas, 'prev_flow': prev_ctrl, 'curr_flow': curr_ctrl
                            })
                        prev_ctrl = curr_ctrl

    elif ext.endswith('.xlsx'):
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active
        headers = [str(cell.value) for cell in ws[1]]
        if 'BG_Gas' in headers:
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[1]: continue
                t_val = row[1]
                dt = t_val if isinstance(t_val, datetime) else parse_mfc_time(str(t_val))
                if dt:
                    events.append({
                        'time': dt, 'bg_gas': str(row[2]), 'bg_flow': float(row[3]),
                        'ctrl_gas': str(row[4]), 'prev_flow': float(row[5]), 'curr_flow': float(row[6])
                    })
    return events

# ================================
# 핵심 메인 프로세스
# ================================
def process_data():
    if not mfc_files and not peak_files:
        messagebox.showwarning("경고", "처리할 파일을 먼저 추가해주세요.")
        return
    if not mfc_files and peak_files:
        messagebox.showerror("에러", "파장(Peak) 파일 단독 처리는 불가합니다.\n왼쪽 목록에 MFC 기준 파일을 반드시 추가해주세요.")
        return

    # [1] MFC 데이터만 있는 경우
    if mfc_files and not peak_files:
        all_events = []
        for f in mfc_files:
            all_events.extend(extract_mfc_events(f))
        
        if not all_events:
            messagebox.showerror("에러", "유효한 MFC 이벤트 데이터가 없습니다.")
            return

        all_events.sort(key=lambda x: x['time'])
        
        # [수정됨] 저장창을 띄우기 전에 전체 데이터의 '가장 첫 시간'을 뽑아서 파일명 생성
        start_time = all_events[0]['time']
        raw_pattern = entry_filename.get()
        prefix = get_formatted_filename(raw_pattern, start_time) if use_time_format_var.get() else raw_pattern
        default_name = f"{prefix}_mfc_setpt.xlsx"

        save_path = filedialog.asksaveasfilename(title="MFC 정리결과 저장", defaultextension=".xlsx", initialfile=default_name, filetypes=[("Excel files", "*.xlsx")])
        if not save_path: return
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['Time', 'BG_Gas', 'BG_Flow', 'Ctrl_Gas', 'Prev_Flow', 'Curr_Flow'])
        for e in all_events:
            ws.append([e['time'].strftime("%y%m%d %H:%M:%S.%f")[:-3], e['bg_gas'], e['bg_flow'], e['ctrl_gas'], e['prev_flow'], e['curr_flow']])
        wb.save(save_path)
        messagebox.showinfo("완료", "MFC 단독 정리가 완료되었습니다.")
        return

    # [2] 올인원 매칭 모드 (Peak 파일 병합)
    output_dir = filedialog.askdirectory(title="파장 데이터 변환 결과를 각각 저장할 '폴더'를 선택하세요")
    if not output_dir: return

    all_events = []
    for f in mfc_files:
        all_events.extend(extract_mfc_events(f))
    all_events.sort(key=lambda x: x['time'])

    success_count = 0
    fail_list = []

    for peak_file in peak_files:
        headers_info = []
        data_rows = []
        is_data = False

        with open(peak_file, 'r', encoding='utf-8') as f:
            for line in f:
                line = line.strip()
                if not line: continue
                if not is_data:
                    if line == "[Data]": is_data = True
                    elif line.startswith("Wavelength"):
                        try:
                            parts = line.split(',')
                            wl = parts[0].split('=')[1]
                            el = parts[1].split('=')[1]
                            headers_info.append(f"{wl} / {el}")
                        except: pass
                else:
                    if not line.startswith("Timestamp"):
                        parts = line.split(',')
                        dt = parse_peak_time(parts[0])
                        if dt: data_rows.append({'dt': dt, 'raw': parts})
        
        if not data_rows:
            fail_list.append(f"{os.path.basename(peak_file)} (데이터 없음)")
            continue

        start_dt = data_rows[0]['dt']
        end_dt = data_rows[-1]['dt']
        
        # [수정됨] 파일별 실제 측정 시작 시간(start_dt)을 기반으로 파일명 생성
        raw_pattern = entry_filename.get()
        prefix = get_formatted_filename(raw_pattern, start_dt) if use_time_format_var.get() else raw_pattern
        
        tolerance = timedelta(minutes=5)
        matched_events = [e for e in all_events if (start_dt - tolerance) <= e['time'] <= (end_dt + tolerance)]
        
        if not matched_events:
            fail_list.append(f"{os.path.basename(peak_file)} (시간대(날짜 포함)가 일치하는 MFC 기록 없음)")
            continue

        blocks_to_write = []
        used_red_indices = set()

        for e in matched_events:
            red_idx = -1
            for i, row in enumerate(data_rows):
                if row['dt'] <= e['time']:
                    red_idx = i
                else:
                    break
            
            if red_idx != -1:
                if abs((e['time'] - data_rows[red_idx]['dt']).total_seconds()) > 60:
                    continue
                
                if red_idx not in used_red_indices:
                    used_red_indices.add(red_idx)
                    start_idx = max(0, red_idx - 10)
                    block = data_rows[start_idx : red_idx + 1]
                    
                    prev_val = int(e['prev_flow']) if e['prev_flow'].is_integer() else e['prev_flow']
                    label = f"{e['ctrl_gas']} {prev_val} SCCM"
                    
                    blocks_to_write.append({
                        'label': label,
                        'block': block
                    })

        if not blocks_to_write:
            fail_list.append(f"{os.path.basename(peak_file)} (매칭된 유효 데이터 없음)")
            continue

        bg_gas = matched_events[0]['bg_gas']
        bg_flow = int(matched_events[0]['bg_flow'])
        ctrl_gas = matched_events[0]['ctrl_gas']
        max_ctrl = max(max(e['prev_flow'], e['curr_flow']) for e in matched_events)
        deltas = [abs(e['curr_flow'] - e['prev_flow']) for e in matched_events if abs(e['curr_flow'] - e['prev_flow']) > 0]
        delta_val = max(set(deltas), key=deltas.count) if deltas else 0
        delta_str = int(delta_val) if delta_val.is_integer() else delta_val

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Peak Data"

        ws['B1'] = f"배경가스 {bg_gas} {bg_flow} sccm + 변화가스 {ctrl_gas} 0~{int(max_ctrl)} sccm(Δ{delta_str} sccm)"
        ws['B1'].font = Font(size=15, bold=True)
        ws.merge_cells('B1:G1')
        ws['B1'].alignment = Alignment(horizontal="left", vertical="center")
        ws['H1'], ws['I1'] = "측정일 :", start_dt.strftime("%y%m%d")
        ws['J1'], ws['K1'] = "수정일 :", datetime.now().strftime("%y%m%d")
        ws['B2'] = "[Data]"
        ws.append(["", "Timestamp", "No"] + headers_info)

        yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
        font_bold = Font(bold=True)
        font_white = Font(color="FFFFFF", bold=True)
        center_align = Alignment(horizontal="center", vertical="center")

        current_row = 4
        for b_dict in blocks_to_write:
            label = b_dict['label']
            block = b_dict['block']
            
            for i, b_row in enumerate(block):
                is_red = (i == len(block) - 1)
                row_data = [label if i == 0 else ""] + b_row['raw']
                ws.append(row_data)

                for col_idx in range(1, len(row_data) + 1):
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.alignment = center_align
                    if is_red:
                        cell.fill = red_fill
                        cell.font = font_white
                    else:
                        cell.fill = yellow_fill
                        if col_idx == 1: cell.font = font_bold
                current_row += 1

        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 25
        ws.freeze_panes = 'B4' 

        # [수정됨] 파일명 덮어쓰기 방지 안전망 추가
        base_save_name = f"{prefix}_peak_{bg_gas}_{int(bg_flow)}_{ctrl_gas}_{int(max_ctrl)}.xlsx"
        save_name = base_save_name
        counter = 1
        while os.path.exists(os.path.join(output_dir, save_name)):
            save_name = base_save_name.replace(".xlsx", f"_{counter}.xlsx")
            counter += 1
            
        wb.save(os.path.join(output_dir, save_name))
        success_count += 1

    msg = f"총 {len(peak_files)}개 중 {success_count}개 변환 성공!\n"
    if fail_list:
        msg += f"\n[ ❌ 실패 내역 ({len(fail_list)}개) ]\n" + "\n".join(fail_list)
        messagebox.showwarning("완료 (일부 실패 존재)", msg)
    else:
        messagebox.showinfo("완료", msg)

# ================================
# GUI 리스트박스 관리 함수
# ================================
def update_lists():
    listbox_mfc.delete(0, tk.END); listbox_peak.delete(0, tk.END)
    for p in mfc_files: listbox_mfc.insert(tk.END, os.path.basename(p))
    for p in peak_files: listbox_peak.insert(tk.END, os.path.basename(p))

def drop_mfc(event):
    for p in root.tk.splitlist(event.data):
        if p not in mfc_files: mfc_files.append(p)
    update_lists()

def drop_peak(event):
    for p in root.tk.splitlist(event.data):
        if p.lower().endswith('.csv') and p not in peak_files: peak_files.append(p)
    update_lists()

def browse_mfc():
    for p in filedialog.askopenfilenames(title="MFC 파일 선택"):
        if p not in mfc_files: mfc_files.append(p)
    update_lists()

def browse_peak():
    for p in filedialog.askopenfilenames(title="Peak 파일 선택", filetypes=[("CSV", "*.csv")]):
        if p not in peak_files: peak_files.append(p)
    update_lists()

def remove_mfc():
    for i in reversed(listbox_mfc.curselection()): del mfc_files[i]
    update_lists()

def remove_peak():
    for i in reversed(listbox_peak.curselection()): del peak_files[i]
    update_lists()

def clear_all():
    mfc_files.clear(); peak_files.clear()
    update_lists()

# ================================
# GUI 화면 구성
# ================================
root = TkinterDnD.Tk()
root.title("MFC & Peak All-in-One Extractor")
root.geometry("700x550")

frame_lists = tk.Frame(root)
frame_lists.pack(pady=10, padx=10, fill="x")

frame_mfc = tk.LabelFrame(frame_lists, text="1. MFC 유량 파일 (Raw 또는 Summary)", padx=5, pady=5)
frame_mfc.pack(side="left", fill="both", expand=True, padx=5)
listbox_mfc = tk.Listbox(frame_mfc, selectmode=tk.EXTENDED, height=10)
listbox_mfc.pack(fill="both", expand=True)
listbox_mfc.drop_target_register(DND_FILES)
listbox_mfc.dnd_bind('<<Drop>>', drop_mfc)

btn_frame_mfc = tk.Frame(frame_mfc)
btn_frame_mfc.pack(fill="x", pady=5)
tk.Button(btn_frame_mfc, text="Browse", command=browse_mfc).pack(side="left", expand=True, fill="x", padx=2)
tk.Button(btn_frame_mfc, text="Remove", command=remove_mfc).pack(side="left", expand=True, fill="x", padx=2)

frame_peak = tk.LabelFrame(frame_lists, text="2. 파장 기록 파일 (Peak CSV)", padx=5, pady=5)
frame_peak.pack(side="left", fill="both", expand=True, padx=5)
listbox_peak = tk.Listbox(frame_peak, selectmode=tk.EXTENDED, height=10)
listbox_peak.pack(fill="both", expand=True)
listbox_peak.drop_target_register(DND_FILES)
listbox_peak.dnd_bind('<<Drop>>', drop_peak)

btn_frame_peak = tk.Frame(frame_peak)
btn_frame_peak.pack(fill="x", pady=5)
tk.Button(btn_frame_peak, text="Browse", command=browse_peak).pack(side="left", expand=True, fill="x", padx=2)
tk.Button(btn_frame_peak, text="Remove", command=remove_peak).pack(side="left", expand=True, fill="x", padx=2)

frame_bot = tk.Frame(root)
frame_bot.pack(fill="x", padx=15)

tk.Button(frame_bot, text="Clear All Files", command=clear_all, width=15).pack(side="right", pady=5)

setting_frame = tk.LabelFrame(root, text="저장 설정 (Peak 출력 시 폴더에 다중 저장됨)", padx=10, pady=10)
setting_frame.pack(pady=5, fill="x", padx=15)

use_time_format_var = tk.BooleanVar(value=True)
tk.Checkbutton(setting_frame, text="파일명에 시간 자동 변환 적용 (yyMMddHHmm 등)", variable=use_time_format_var).grid(row=0, column=0, columnspan=2, sticky="w")
tk.Label(setting_frame, text="파일명 접두사:").grid(row=1, column=0, sticky="w", pady=5)
entry_filename = tk.Entry(setting_frame, width=30)
entry_filename.insert(0, "yyMMddHHmm")
entry_filename.grid(row=1, column=1, sticky="w", padx=5)

tk.Button(root, text="Process & Save All", command=process_data, height=2, bg="#4CAF50", fg="white", font=("", 12, "bold")).pack(fill="x", padx=15, pady=10)

root.mainloop()